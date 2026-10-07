"""② 記入テンプレート(LINE貼り付け用テキスト+Excel)を spec/items.yaml から生成する。

使い方: python scripts/build_templates.py [--skip-theme 6 ...]
  → templates/line/00_案内.txt, 01_料金セット・延長メニュー.txt …
  → templates/excel/初期設定_記入シート.xlsx
"""

from __future__ import annotations

import argparse
import base64
import datetime
import io
import json
import re
import sys
import zipfile
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
from validate_items import load_spec  # noqa: E402

ROOT = Path(__file__).resolve().parent.parent
LINE_DIR = ROOT / "templates" / "line"
XLSX = ROOT / "templates" / "excel" / "初期設定_記入シート.xlsx"
# Googleスプレッドシート版の URL(Drive へアップロード後に記入。無ければページにボタンを出さない)
SHEETS_YAML = ROOT / "templates" / "google_sheets.yaml"
PREVIEW_TEMPLATE = ROOT / "templates" / "kit" / "preview.html"
PREVIEW = ROOT / "templates" / "preview" / "index.html"

INPUT_ROWS = 30  # 記入例の下に用意する空行の数

# LINEで最初に送る案内文(6章の文章ルールに合わせて改行済み)
GUIDE = """\
アンケートのご回答、
ありがとうございました。

次に、下のひな形をLINEで送ります。
テーマごとに1通ずつ、
記入して返信していただけると助かります。

「例)」の行は消して、
お店の内容に書き換えてください。
写真でOKのものは、撮って送るだけで大丈夫です。

わからない所は、空欄のままで構いません。
訪問時に一緒に確認します。

Excelの方が書きやすい場合は、
添付の記入シートをお使いください。"""

CLOSING = """\
すべて送っていただいたら、
こちらでPOSへの入力を進めます。

入力が終わったら、
訪問日の前にご連絡します。
どうぞよろしくお願いいたします。"""


def template_items(spec: dict, theme_id: int) -> list[dict]:
    return [
        it for it in spec["items"]
        if it["route"] == "template" and it["theme"] == theme_id and "merged_into" not in it
    ]


def themes(spec: dict, skip: set[int] = frozenset()) -> list[dict]:
    return [t for t in spec["themes"] if t["id"] not in skip]


# ── LINE ────────────────────────────────────────────


def line_text(spec: dict, theme: dict) -> str:
    items = template_items(spec, theme["id"])
    lines = [f"【{theme['title']}】"]
    lines += [f"※{n}" for n in theme.get("line_notes", [])]
    for it in items:
        if len(items) > 1:
            lines.append(f"■{it['label']}")
        fmt = it["line"]["format"]
        lines.append("・" + "/".join(fmt))
        for ex in it["line"]["examples"]:
            lines.append("・例)" + "/".join(str(v) for v in ex))
    lines += theme.get("line_after", [])
    return "\n".join(lines)


def line_messages(spec: dict, skip: set[int] = frozenset()) -> list[tuple[str, str]]:
    """[(ファイル名, 本文)] を送る順に返す。"""
    msgs = [("00_案内", GUIDE)]
    for t in themes(spec, skip):
        msgs.append((f"{t['id']:02d}_{t['title']}", line_text(spec, t)))
    msgs.append(("99_最後に", CLOSING))
    return msgs


# ── Excel ───────────────────────────────────────────

ACCENT = "673AB7"
HEAD_FILL = PatternFill("solid", fgColor="EDE4F8")
EXAMPLE_FONT = Font(color="9AA0A6", italic=True)
TITLE_FONT = Font(bold=True, size=14, color=ACCENT)
BLOCK_FONT = Font(bold=True, size=12)
NOTE_FONT = Font(color="B3261E")
THIN = Side(style="thin", color="DADCE0")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _validation(cell_type, first: str, last: str) -> DataValidation | None:
    rng = f"{first}:{last}"
    if isinstance(cell_type, list):
        dv = DataValidation(type="list", formula1='"' + ",".join(cell_type) + '"', allow_blank=True)
        dv.error, dv.errorTitle = "リストから選んでください", "入力エラー"
    elif cell_type == "int":
        dv = DataValidation(type="whole", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
        dv.error, dv.errorTitle = "数字のみで入力してください(カンマ不要)", "入力エラー"
    elif cell_type == "num":
        dv = DataValidation(type="decimal", operator="greaterThanOrEqual", formula1="0", allow_blank=True)
        dv.error, dv.errorTitle = "0以上の数字で入力してください", "入力エラー"
    else:
        return None
    dv.showErrorMessage = True
    dv.add(rng)
    return dv


def sheet_rows(spec: dict, theme: dict, input_rows: int = INPUT_ROWS, standalone: bool = False) -> list[dict]:
    """1シート分の行を、種類(title / store / note / block / head / example / input / blank)付きで返す。

    standalone … 店舗に配るファイルのとき。店舗名の欄と、返送の案内を付ける。
    """
    rows = [{"kind": "title", "cells": [theme["title"]]}]
    if standalone:
        rows.append({"kind": "store", "cells": ["店舗名", None]})
    notes = ["灰色の行は記入例です。消して上書きするか、下の空欄に書いてください。"] + theme.get("excel_notes", [])
    if standalone:
        notes.append("書き終わったら、このファイル(またはシートのリンク)をLINEで送ってください。")
    rows += [{"kind": "note", "cells": ["※" + n], "warn": "個人情報" in n} for n in notes]
    for it in template_items(spec, theme["id"]):
        rows.append({"kind": "blank", "cells": []})
        rows.append({"kind": "block", "cells": [it["label"]]})
        rows.append({"kind": "head", "cells": list(it["columns"])})
        rows += [{"kind": "example", "cells": [None if v == "" else v for v in ex]} for ex in it["examples"]]
        dropdown = [c if isinstance(c, list) else None for c in it["cells"]]
        numeric = [c in ("int", "num") for c in it["cells"]]
        rows += [{"kind": "input", "cells": [], "width": len(it["columns"]), "cells_type": it["cells"],
                  "dropdown": dropdown, "numeric": numeric} for _ in range(input_rows)]
    return rows


def build_workbook(spec: dict, skip: set[int] = frozenset(), standalone: bool = False) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)
    for t in themes(spec, skip):
        ws = wb.create_sheet(t["sheet"])
        ws.sheet_view.showGridLines = False
        rows = sheet_rows(spec, t, standalone=standalone)
        widths: dict[int, int] = {}
        r = 0
        while r < len(rows):
            row, n = rows[r], r + 1
            kind = row["kind"]
            if kind == "input":
                # 同じ表の入力行をまとめて、罫線と入力規則を付ける
                last = r
                while last + 1 < len(rows) and rows[last + 1]["kind"] == "input":
                    last += 1
                for rr in range(n, last + 2):
                    for c in range(1, row["width"] + 1):
                        ws.cell(row=rr, column=c).border = BOX
                for c, ct in enumerate(row["cells_type"], start=1):
                    col = get_column_letter(c)
                    dv = _validation(ct, f"{col}{n}", f"{col}{last + 1}")
                    if dv:
                        ws.add_data_validation(dv)
                r = last + 1
                continue
            for c, v in enumerate(row["cells"], start=1):
                cell = ws.cell(row=n, column=c, value=v)
                if kind == "title":
                    cell.font = TITLE_FONT
                elif kind == "store":
                    cell.font = BLOCK_FONT if c == 1 else Font()
                    if c == 2:
                        cell.border = BOX
                elif kind == "note":
                    cell.font = NOTE_FONT if row.get("warn") else Font(color="5F6368")
                elif kind == "block":
                    cell.font = BLOCK_FONT
                elif kind == "head":
                    cell.font, cell.fill, cell.border = Font(bold=True), HEAD_FILL, BOX
                    cell.alignment = Alignment(horizontal="center")
                    widths[c] = max(widths.get(c, 0), len(str(v)) * 2 + 4)
                elif kind == "example":
                    cell.font, cell.border = EXAMPLE_FONT, BOX
            r += 1
        for c, w in widths.items():
            ws.column_dimensions[get_column_letter(c)].width = max(12, min(w, 28))
    return wb


FIXED_TIME = datetime.datetime(2026, 1, 1)


def _xlsx_bytes(wb: Workbook) -> bytes:
    """作成時刻を固定して書き出す。同じ入力なら、いつ作っても同じバイト列になる。"""
    wb.properties.created = wb.properties.modified = FIXED_TIME
    raw = io.BytesIO()
    wb.save(raw)
    out = io.BytesIO()
    with zipfile.ZipFile(raw) as src, zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            fixed = zipfile.ZipInfo(info.filename, FIXED_TIME.timetuple()[:6])
            fixed.compress_type = zipfile.ZIP_DEFLATED
            fixed.external_attr = info.external_attr
            data = src.read(info.filename)
            if info.filename == "docProps/core.xml":
                # openpyxl は保存の瞬間に更新日時を入れるため、固定値へ置き換える
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*", rb"\g<1>2026-01-01T00:00:00Z", data)
            dst.writestr(fixed, data)
    return out.getvalue()


def spreadsheet_url() -> str | None:
    """Googleスプレッドシート版(全シートをまとめた1ファイル)の URL。"""
    if not SHEETS_YAML.exists():
        return None
    data = yaml.safe_load(SHEETS_YAML.read_text(encoding="utf-8")) or {}
    return data.get("spreadsheet") or None


def build_preview(spec: dict, xlsx_bytes: bytes) -> str:
    """LINE・Excel・スプレッドシートを確認するページ(templates/preview/index.html)。"""
    optional = {f"{t['id']:02d}_{t['title']}" for t in spec["themes"] if t.get("optional")}
    titles = {"00_案内": "ご案内", "99_最後に": "最後に"}
    messages = [
        {"title": titles.get(name, name.split("_", 1)[1]), "body": body, "optional": name in optional}
        for name, body in line_messages(spec)
    ]
    sheets = [{"name": t["sheet"], "title": t["title"], "rows": [
        {k: v for k, v in row.items() if k != "cells_type"}
        for row in sheet_rows(spec, t, input_rows=3, standalone=True)
    ]} for t in themes(spec)]
    data = {
        "messages": messages,
        "sheets": sheets,
        "xlsx": {"filename": XLSX.name, "base64": base64.b64encode(xlsx_bytes).decode()},
        "spreadsheetUrl": spreadsheet_url(),
    }
    js = json.dumps(data, ensure_ascii=False).replace("</", "<\\/")
    return PREVIEW_TEMPLATE.read_text(encoding="utf-8").replace("/*__DATA_JSON__*/null", js)


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--skip-theme", type=int, nargs="*", default=[], help="送らないテーマ(①で「使わない」)")
    args = ap.parse_args()
    spec, skip = load_spec(), set(args.skip_theme)
    LINE_DIR.mkdir(parents=True, exist_ok=True)
    for old in LINE_DIR.glob("*.txt"):
        old.unlink()
    for name, body in line_messages(spec, skip):
        (LINE_DIR / f"{name}.txt").write_text(body + "\n", encoding="utf-8")
    # 記入シートは1ファイルに全種類(1種類1シート)。Googleスプレッドシート版はこれを変換したもの
    XLSX.parent.mkdir(parents=True, exist_ok=True)
    XLSX.write_bytes(_xlsx_bytes(build_workbook(spec, skip, standalone=True)))
    # 確認ページは、すべてのテーマを載せる
    PREVIEW.parent.mkdir(parents=True, exist_ok=True)
    PREVIEW.write_text(build_preview(spec, _xlsx_bytes(build_workbook(spec, standalone=True))), encoding="utf-8")
    print(f"wrote {LINE_DIR.relative_to(ROOT)}/*.txt, {XLSX.relative_to(ROOT)} and {PREVIEW.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
